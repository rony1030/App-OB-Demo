from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "output" / "documents"
OUT_DIR.mkdir(parents=True, exist_ok=True)
OUT_PATH = OUT_DIR / "Guia_entrega_Elements_Developer.docx"

INK = RGBColor(26, 26, 26)
MUTED = RGBColor(90, 86, 80)
BRONZE = RGBColor(169, 129, 98)


def set_run_font(run, name="Arial", size=8.3, bold=False, color=INK):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color


def set_columns(section, count=2, space_twips=420):
    sect_pr = section._sectPr
    cols = sect_pr.xpath("./w:cols")
    if cols:
        cols_el = cols[0]
    else:
        cols_el = OxmlElement("w:cols")
        sect_pr.append(cols_el)
    cols_el.set(qn("w:num"), str(count))
    cols_el.set(qn("w:space"), str(space_twips))


def compact(paragraph, before=0, after=0, line=1.0, keep=False):
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line
    fmt.keep_with_next = keep


def section_heading(doc, number, title):
    p = doc.add_paragraph()
    compact(p, before=5, after=2, line=1.0, keep=True)
    r = p.add_run(f"{number}  {title}")
    set_run_font(r, size=9.4, bold=True, color=INK)
    return p


def item(doc, text, checked=False):
    p = doc.add_paragraph()
    compact(p, before=0, after=1.1, line=1.0)
    p.paragraph_format.left_indent = Inches(0.04)
    marker = "[x]" if checked else "[ ]"
    r = p.add_run(f"{marker}  {text}")
    set_run_font(r, size=7.7, color=INK)
    return p


def note(doc, lead, text):
    p = doc.add_paragraph()
    compact(p, before=2, after=2, line=1.05)
    r = p.add_run(lead)
    set_run_font(r, size=7.8, bold=True, color=BRONZE)
    r = p.add_run(text)
    set_run_font(r, size=7.8, color=MUTED)
    return p


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(0.42)
section.bottom_margin = Inches(0.38)
section.left_margin = Inches(0.5)
section.right_margin = Inches(0.5)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Arial"
normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
normal.font.size = Pt(8.3)
normal.font.color.rgb = INK

title = doc.add_paragraph(style="Title")
title.alignment = WD_ALIGN_PARAGRAPH.LEFT
compact(title, after=3, line=0.95)
tr = title.add_run("Guía de entrega para Elements Developer")
set_run_font(tr, size=21, bold=True, color=INK)

subtitle = doc.add_paragraph()
compact(subtitle, after=5, line=1.05)
sr = subtitle.add_run("Landing pública, inventario y presencia dentro de la plataforma OB Brokers")
set_run_font(sr, size=9.4, bold=True, color=BRONZE)

intro = doc.add_paragraph()
compact(intro, after=5, line=1.1)
ir = intro.add_run(
    "Elements se incorporará como proyecto, igual que Ciprés, bajo la desarrolladora Elements Developer. "
    "Entrega un solo archivo ZIP con la estructura indicada. Si la información publicada en elements.com.do sigue vigente, marca CONFIRMADO y autoriza su extracción; no es necesario duplicarla."
)
set_run_font(ir, size=8.4, color=MUTED)

set_columns(section, 2, 430)

section_heading(doc, "01", "Identidad y marca")
item(doc, "Logo horizontal y símbolo en SVG; respaldo PNG transparente de 2000 px o más")
item(doc, "Manual de marca, colores oficiales y tipografías con sus licencias")
item(doc, "Nombre comercial confirmado: ELEMENTS y desarrolladora: Elements Developer")
item(doc, "Versiones del logo para fondo claro, oscuro y fotografía")

section_heading(doc, "02", "Imágenes y movimiento")
item(doc, "1 hero horizontal de al menos 2400 x 1600 px, limpio y sin texto incrustado")
item(doc, "8 a 15 renders exteriores e interiores en alta resolución")
item(doc, "2 a 4 imágenes verticales para móvil y revelaciones editoriales")
item(doc, "Planos, master plan, mapa de ubicación y fotografías de amenidades")
item(doc, "Video hero opcional de 10 a 30 s, MP4 WebM, 1080p o superior, sin música obligatoria")
item(doc, "Autorización para reutilizar los recursos públicos de elements.com.do si faltan originales")

section_heading(doc, "03", "Información comercial vigente")
item(doc, "Nombre oficial, fase, estado de venta, ubicación y fecha estimada de entrega")
item(doc, "Tipologías, metrajes interiores y exteriores, capacidad y equipamiento incluido")
item(doc, "Precio desde, moneda, reserva, separación y planes de pago")
item(doc, "Inventario por unidad: código, modelo, precio, estado y disponibilidad")
item(doc, "Amenidades, atributos sostenibles, paneles solares, picuzzi y upgrades")
item(doc, "Comisión de broker, incentivos, pronto pago y vigencias autorizadas")

section_heading(doc, "04", "Ventas, contacto y conversión")
item(doc, "WhatsApp con código de país, teléfono, correo y nombre del responsable comercial")
item(doc, "Destinatarios del formulario y tiempo esperado de respuesta")
item(doc, "Instagram y demás enlaces oficiales")
item(doc, "CTA principal: reservar unidad, solicitar disponibilidad, invertir u otro")
item(doc, "Ruta o dominio final y acceso deseado desde el portal")

section_heading(doc, "05", "Documentos y respaldo")
item(doc, "Brochure, lista de precios, inventario y esquema de pagos vigentes")
item(doc, "Memoria de calidades, especificaciones y planos aprobados")
item(doc, "Aviso legal sobre renders, privacidad, términos y cookies")
item(doc, "Fuente del inventario: Supabase, Google Sheets o carga manual")

section_heading(doc, "06", "Datos detectados en la web para confirmar")
item(doc, "ELEMENTS Eco Residences & Resort; Fase 0; 12 Eco Suites")
item(doc, "Modelos publicados: Eco Suite, Eco Deluxe y Eco Royal")
item(doc, "Eco Suite desde US$95,000 y reserva publicada de US$500")
item(doc, "Planes A y B; descuento publicado de 5% por pronto pago bajo condición")
item(doc, "Concepto de marca: Tierra, Agua, Aire y Fuego; enfoque sostenible y natural")
note(doc, "Acción: ", "escribe CONFIRMADO, CORREGIR o RETIRAR junto a cada dato antes de publicarlo.")

section_heading(doc, "07", "Estructura del archivo ZIP")
for folder in [
    "01_Marca",
    "02_Renders_Horizontales",
    "03_Renders_Verticales",
    "04_Videos",
    "05_Planos_y_Ubicacion",
    "06_Inventario_y_Precios",
    "07_Documentos_Legales",
    "08_Contactos_y_Enlaces",
]:
    item(doc, folder)

section_heading(doc, "08", "Decisiones rápidas")
item(doc, "Idiomas: español / inglés / francés")
item(doc, "Inventario: manual / Google Sheets / Supabase")
item(doc, "Página: nueva ruta / subdominio / dominio propio")
item(doc, "Flujo visual recomendado: comp primero y luego desarrollo")

note(doc, "Entrega mínima para comenzar: ", "logo, 6 renders originales, datos comerciales confirmados, contacto de ventas y autorización de uso de activos.")
note(doc, "Referencia visual: ", "arquitectura editorial de Vistal y revelaciones cinemáticas inspiradas en Gisou, adaptadas a la identidad de Elements.")

footer = section.footer
fp = footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
compact(fp, before=0, after=0)
fr = fp.add_run("OB Brokers Team  |  Preparación de contenido para Elements Developer  |  Septiembre 2026")
set_run_font(fr, size=6.8, color=MUTED)

doc.save(OUT_PATH)
print(OUT_PATH)
